"""Спільні функції: Prozorro, Telegram, документи, Claude API."""
import os, re, io, time, html, zipfile, base64, logging
import requests

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("tender")

API = "https://public-api.prozorro.gov.ua/api/2.5"
TG_TOKEN = os.environ.get("TG_TOKEN", "").strip()
TG_CHAT = os.environ.get("TG_CHAT_ID", "").strip().strip('"')
ANTHROPIC_KEY = os.environ.get("ANTHROPIC_API_KEY", "").strip()
MODEL_FULL = os.getenv("MODEL_FULL", "claude-sonnet-5")
MODEL_SHORT = os.getenv("MODEL_SHORT", "claude-haiku-4-5-20251001")
PRICES = {"claude-sonnet-5": (2, 10), "claude-haiku-4-5-20251001": (1, 5)}  # $ за 1 млн токенів (вхід, вихід)

S = requests.Session()
S.headers["User-Agent"] = "tender-monitor/2.0"


# ---------------- HTTP ----------------
def get(url, params=None):
    for i in range(5):
        try:
            r = S.get(url, params=params, timeout=60)
            if r.status_code == 429 or r.status_code >= 500:
                raise requests.HTTPError(f"HTTP {r.status_code}")
            r.raise_for_status()
            return r.json()
        except Exception as e:
            log.warning("GET %s: %s (спроба %d)", url, e, i + 1)
            time.sleep(2 ** i)
    raise RuntimeError(f"API недоступне: {url}")


def download(url, limit=30_000_000):
    for i in range(3):
        try:
            with S.get(url, timeout=120, stream=True) as r:
                r.raise_for_status()
                buf = io.BytesIO()
                for chunk in r.iter_content(1 << 16):
                    buf.write(chunk)
                    if buf.tell() > limit:
                        raise ValueError("файл завеликий")
                return buf.getvalue()
        except ValueError:
            raise
        except Exception as e:
            log.warning("download %s: %s", url[:80], e)
            time.sleep(3 * (i + 1))
    raise RuntimeError("не вдалося завантажити")


# ---------------- Telegram ----------------
def tg_call(method, data=None, files=None):
    url = f"https://api.telegram.org/bot{TG_TOKEN}/{method}"
    for _ in range(3):
        r = requests.post(url, data=data, files=files, timeout=120) if files else \
            requests.post(url, json=data or {}, timeout=60)
        j = r.json()
        if r.status_code == 429:
            time.sleep(j.get("parameters", {}).get("retry_after", 5)); continue
        if j.get("ok"):
            return j.get("result")
        raise RuntimeError(f"Telegram {method}: {j.get('description')} (chat_id={TG_CHAT!r})")
    raise RuntimeError(f"Telegram {method}: перевищено ліміт")


def tg_send(text, buttons=None, chat=None):
    data = {"chat_id": chat or TG_CHAT, "text": text[:4096], "parse_mode": "HTML",
            "disable_web_page_preview": True}
    if buttons:
        data["reply_markup"] = {"inline_keyboard": buttons}
    try:
        tg_call("sendMessage", data)
    except RuntimeError as e:
        if "parse entities" not in str(e):
            raise
        data.pop("parse_mode")
        data["text"] = html.unescape(re.sub(r"<[^>]+>", "", text))[:4096]
        tg_call("sendMessage", data)
    time.sleep(1.1)


def tg_send_file(name, content, caption="", chat=None):
    tg_call("sendDocument", {"chat_id": chat or TG_CHAT, "caption": caption[:1000]},
            files={"document": (name, content)})
    time.sleep(1.1)


# ---------------- Prozorro ----------------
UA_RE = re.compile(r"UA-\d{4}-\d{2}-\d{2}-\d{6}-[a-z]")
HEX_RE = re.compile(r"\b[0-9a-f]{32}\b")


def tender(tid):
    return get(f"{API}/tenders/{tid}")["data"]


def resolve_id(text):
    """Внутрішній id тендеру з тексту: посилання, UA-номер або 32-символьний id."""
    m = HEX_RE.search(text)
    if m:
        return m.group(0)
    m = UA_RE.search(text)
    if not m:
        return None
    ua = m.group(0)
    try:
        r = S.get(f"https://prozorro.gov.ua/api/tenders/{ua}/summary", timeout=30)
        cands = list(dict.fromkeys(HEX_RE.findall(r.text)))[:8]
    except Exception as e:
        log.warning("summary %s: %s", ua, e); cands = []
    for c in cands:
        try:
            if tender(c).get("tenderID") == ua:
                return c
        except Exception:
            pass
    return None


def dt(s):
    from datetime import datetime
    return datetime.fromisoformat(s).strftime("%d.%m.%Y %H:%M") if s else "—"


def meta(t):
    """Стислі метадані тендеру для аналізу."""
    v, pe = t.get("value") or {}, t.get("procuringEntity") or {}
    L = [f"Номер: {t.get('tenderID')}", f"Назва: {t.get('title')}", f"Опис: {t.get('description') or '—'}",
         f"Замовник: {pe.get('name')} (ЄДРПОУ {(pe.get('identifier') or {}).get('id')}), "
         f"{(pe.get('address') or {}).get('region')}, {(pe.get('address') or {}).get('locality')}",
         f"Очікувана вартість: {v.get('amount')} {v.get('currency')} "
         f"{'з ПДВ' if v.get('valueAddedTaxIncluded') else 'без ПДВ'}",
         f"Процедура: {t.get('procurementMethodType')}; категорія: {t.get('mainProcurementCategory')}; статус: {t.get('status')}",
         f"Уточнення до: {dt((t.get('enquiryPeriod') or {}).get('endDate'))}",
         f"Подання пропозицій до: {dt((t.get('tenderPeriod') or {}).get('endDate'))}",
         f"Забезпечення пропозиції: {t.get('guarantee') or 'не вимагається/не вказано'}"]
    for l in t.get("lots") or []:
        L.append(f"Лот: {l.get('title')} — {(l.get('value') or {}).get('amount')} грн; забезпечення: {l.get('guarantee')}")
    L.append("Позиції:")
    for i in (t.get("items") or [])[:60]:
        a = (i.get("deliveryAddress") or {})
        L.append(f"- {i.get('description')} | {i.get('quantity')} {(i.get('unit') or {}).get('name', '')} | "
                 f"ДК {(i.get('classification') or {}).get('id')} | до {dt((i.get('deliveryDate') or {}).get('endDate'))} | "
                 f"{a.get('region', '')} {a.get('locality', '')} {a.get('streetAddress', '')}")
    if t.get("milestones"):
        L.append("Умови оплати (milestones):")
        for m in t["milestones"]:
            L.append(f"- {m.get('title')} / {m.get('code')}: {m.get('percentage')}%, {(m.get('duration') or {}).get('days')} "
                     f"{(m.get('duration') or {}).get('type')} дн. {m.get('description') or ''}")
    crit = [c for c in t.get("criteria") or [] if not (c.get("classification") or {}).get("id", "").startswith("CRITERION.EXCLUSION")]
    if crit:
        L.append("Критерії (кваліфікаційні та інші):")
        for c in crit:
            reqs = "; ".join(r.get("title", "") for g in c.get("requirementGroups") or [] for r in g.get("requirements") or [])
            L.append(f"- {c.get('title')}: {(c.get('description') or '')[:400]} → {reqs[:600]}")
    if t.get("questions"):
        L.append("Питання та відповіді:")
        for q in t["questions"][:30]:
            L.append(f"- П: {(q.get('title') or '')} {(q.get('description') or '')[:400]}\n  В: {(q.get('answer') or 'без відповіді')[:600]}")
    return "\n".join(L)[:30_000]


# ---------------- Документи ----------------
PRIORITY = {"tenderNotice": 0, "biddingDocuments": 1, "technicalSpecifications": 2, "eligibilityCriteria": 3,
            "evaluationCriteria": 3, "contractProforma": 4, "contractDraft": 4, "billOfQuantity": 5,
            "clarifications": 6, "riskProvisions": 6}
SKIP_EXT = (".p7s", ".sig", ".asic", ".asice", ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".dwg", ".dxf", ".rar", ".7z")


def latest_docs(t):
    latest = {}
    for d in t.get("documents") or []:
        k = d.get("id")
        if k not in latest or (d.get("dateModified") or "") >= (latest[k].get("dateModified") or ""):
            latest[k] = d
    docs = [d for d in latest.values() if d.get("url")]
    docs.sort(key=lambda d: (PRIORITY.get(d.get("documentType"), 7), d.get("title", "")))
    return docs


def extract(name, data, depth=0):
    """Повертає список {'name','text'} | {'name','pdf'} | {'name','skip'}."""
    n = name.lower()
    try:
        if n.endswith(SKIP_EXT):
            return [{"name": name, "skip": "формат не аналізується"}]
        if n.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for tb in d.tables:
                for row in tb.rows:
                    cells = list(dict.fromkeys(c.text.strip() for c in row.cells))
                    parts.append(" | ".join(cells))
            return [{"name": name, "text": "\n".join(parts)}]
        if n.endswith((".xlsx", ".xlsm")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"--- Аркуш: {ws.title}")
                for k, row in enumerate(ws.iter_rows(values_only=True)):
                    if k > 3000: break
                    vals = [str(v) for v in row if v not in (None, "")]
                    if vals: out.append("\t".join(vals))
            return [{"name": name, "text": "\n".join(out)}]
        if n.endswith(".xls"):
            import xlrd
            wb = xlrd.open_workbook(file_contents=data)
            out = []
            for sh in wb.sheets():
                out.append(f"--- Аркуш: {sh.name}")
                for r in range(min(sh.nrows, 3000)):
                    vals = [str(v) for v in sh.row_values(r) if v not in ("", None)]
                    if vals: out.append("\t".join(vals))
            return [{"name": name, "text": "\n".join(out)}]
        if n.endswith(".pdf"):
            from pypdf import PdfReader
            txt = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
            if len(txt.strip()) < 300:
                return [{"name": name, "pdf": data}]          # скан — передамо Claude як PDF
            return [{"name": name, "text": txt}]
        if n.endswith((".txt", ".csv", ".htm", ".html", ".xml", ".json")):
            for enc in ("utf-8", "cp1251"):
                try: return [{"name": name, "text": data.decode(enc)}]
                except UnicodeDecodeError: pass
        if n.endswith(".zip") and depth < 2:
            out = []
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for zi in z.infolist()[:40]:
                    if zi.is_dir(): continue
                    nm = zi.filename
                    try: nm = nm.encode("cp437").decode("cp866")
                    except Exception: pass
                    out += extract(f"{name}/{nm}", z.read(zi), depth + 1)
            return out
        return [{"name": name, "skip": "формат не підтримується (.doc/.rtf тощо)"}]
    except Exception as e:
        return [{"name": name, "skip": f"помилка читання: {e}"}]


def collect_docs(t, budget=250_000, per_doc=90_000, max_scans=2):
    parts, skipped, used, scans = [], [], 0, 0
    for d in latest_docs(t):
        title = d.get("title", "document")
        if title.lower().endswith(SKIP_EXT) or d.get("format") == "application/pkcs7-signature":
            continue
        if used >= budget:
            skipped.append(f"{title} (перевищено ліміт обсягу)"); continue
        try:
            data = download(d["url"])
        except Exception as e:
            skipped.append(f"{title} ({e})"); continue
        for it in extract(title, data):
            if "text" in it and it["text"].strip():
                txt = it["text"][:min(per_doc, budget - used)]
                if len(txt) < len(it["text"]):
                    txt += "\n[...документ обрізано через ліміт обсягу...]"
                parts.append({"name": it["name"], "text": txt}); used += len(txt)
            elif "pdf" in it and scans < max_scans and len(it["pdf"]) < 15_000_000:
                parts.append(it); scans += 1
            elif "skip" in it:
                skipped.append(f"{it['name']} ({it['skip']})")
            elif "pdf" in it:
                skipped.append(f"{it['name']} (скан, перевищено ліміт сканів)")
    return parts, skipped


# ---------------- Claude API ----------------
def claude(system, blocks, model, max_tokens):
    if not ANTHROPIC_KEY:
        raise RuntimeError("не задано секрет ANTHROPIC_API_KEY")
    body = {"model": model, "max_tokens": max_tokens, "system": system,
            "messages": [{"role": "user", "content": blocks}]}
    hdr = {"x-api-key": ANTHROPIC_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"}
    for i in range(4):
        r = requests.post("https://api.anthropic.com/v1/messages", headers=hdr, json=body, timeout=900)
        if r.status_code in (429, 500, 502, 503, 529):
            log.warning("Claude API %s, повтор", r.status_code); time.sleep(30 * (i + 1)); continue
        if not r.ok:
            raise RuntimeError(f"Claude API {r.status_code}: {r.text[:300]}")
        j = r.json()
        text = "".join(b.get("text", "") for b in j.get("content", []) if b.get("type") == "text")
        u = j.get("usage", {})
        pin, pout = PRICES.get(model, (0, 0))
        cost = (u.get("input_tokens", 0) * pin + u.get("output_tokens", 0) * pout) / 1e6
        log.info("Claude %s: вхід %s, вихід %s токенів, ~$%.3f", model, u.get("input_tokens"), u.get("output_tokens"), cost)
        return text, cost
    raise RuntimeError("Claude API перевантажений, спробуйте пізніше")


def blocks_for(t, parts, skipped):
    txt = [f"МЕТАДАНІ ТЕНДЕРУ (Prozorro):\n{meta(t)}\n"]
    pdfs = []
    for p in parts:
        if "text" in p:
            txt.append(f"\n=== ДОКУМЕНТ: {p['name']} ===\n{p['text']}")
        else:
            pdfs.append({"type": "document", "title": p["name"][:200],
                         "source": {"type": "base64", "media_type": "application/pdf",
                                    "data": base64.b64encode(p["pdf"]).decode()}})
    if skipped:
        txt.append("\nНЕ ВДАЛОСЯ ПРОАНАЛІЗУВАТИ (врахуй у висновках): " + "; ".join(skipped))
    return pdfs + [{"type": "text", "text": "\n".join(txt)}]
